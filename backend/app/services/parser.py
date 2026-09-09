"""
ResumeIQ - Resume Parser.

Extracts readable text from supported resume formats.

Supported formats:
    - PDF
    - DOCX
    - TXT
"""

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
}


class ResumeParserError(Exception):
    """Raised when a resume cannot be parsed."""


def extract_text(
    file_content: bytes,
    filename: str,
) -> str:
    """
    Extract text from a supported resume file.

    Args:
        file_content: Raw uploaded file bytes.
        filename: Original uploaded filename.

    Returns:
        Extracted resume text.

    Raises:
        ResumeParserError: If the file type is unsupported
        or the file cannot be parsed.
    """

    if not file_content:
        raise ResumeParserError(
            "The uploaded file is empty."
        )

    if not filename or not filename.strip():
        raise ResumeParserError(
            "A valid filename is required."
        )

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParserError(
            f"Unsupported file type: {extension or 'unknown'}."
        )

    try:
        if extension == ".pdf":
            return _extract_pdf(file_content)

        if extension == ".docx":
            return _extract_docx(file_content)

        if extension == ".txt":
            return _extract_txt(file_content)

    except ResumeParserError:
        raise

    except Exception as exc:
        raise ResumeParserError(
            "Unable to read the uploaded resume."
        ) from exc

    raise ResumeParserError(
        "Unable to determine the resume file format."
    )


def _extract_pdf(file_content: bytes) -> str:
    """Extract text from a PDF document."""

    try:
        reader = PdfReader(
            BytesIO(file_content)
        )
    except Exception as exc:
        raise ResumeParserError(
            "The uploaded PDF could not be opened."
        ) from exc

    if not reader.pages:
        raise ResumeParserError(
            "The uploaded PDF contains no pages."
        )

    pages = []

    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            raise ResumeParserError(
                "Unable to extract text from the PDF."
            ) from exc

        if page_text.strip():
            pages.append(page_text.strip())

    return "\n\n".join(pages).strip()


def _extract_docx(file_content: bytes) -> str:
    """Extract text from a DOCX document."""

    try:
        document = Document(
            BytesIO(file_content)
        )
    except Exception as exc:
        raise ResumeParserError(
            "The uploaded DOCX file could not be opened."
        ) from exc

    sections = []

    # -------------------------------------------------------------
    # Paragraphs
    # -------------------------------------------------------------

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # -------------------------------------------------------------
    # Tables
    #
    # Resume information is frequently stored inside tables,
    # especially in modern Word resume templates.
    # -------------------------------------------------------------

    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                sections.append(" | ".join(cells))

    return "\n".join(sections).strip()


def _extract_txt(file_content: bytes) -> str:
    """Extract text from a plain-text resume."""

    try:
        return file_content.decode(
            "utf-8",
            errors="replace",
        ).strip()

    except Exception as exc:
        raise ResumeParserError(
            "Unable to read the text resume."
        ) from exc